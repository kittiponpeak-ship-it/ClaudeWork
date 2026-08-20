#!/usr/bin/env python3
"""Build an interactive, Google-Maps-styled customer map per month from the source workbooks.

Usage:
    pip install openpyxl
    python3 tools/build_customer_maps.py

Reads data/*.xlsx (sheets "Customers with Coordinates", "No Match", "Notes") and
writes one self-contained HTML file per month into maps/. Leaflet is inlined, so
the output opens straight from disk; only the base-map tiles need a connection.
"""

import json
import re
import unicodedata
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
TEMPLATE = ROOT / "tools" / "map_template.html"
VENDOR = ROOT / "tools" / "vendor"
OUT_DIR = ROOT / "maps"

SOURCES = [
    {
        "xlsx": ROOT / "data" / "January_2026_Customers_with_Coordinates.xlsx",
        "slug": "January_2026",
        "title": "Retail Customer Mapping — January 2026",
        "month": "January 2026",
    },
    {
        "xlsx": ROOT / "data" / "March_2026_Customers_with_Coordinates.xlsx",
        "slug": "March_2026",
        "title": "Retail Customer Mapping — March 2026",
        "month": "March 2026",
    },
]

# Google Maps marker / brand colours, handed out by descending turnover
GOOGLE_PALETTE = [
    "#EA4335",  # red
    "#4285F4",  # blue
    "#34A853",  # green
    "#FBBC04",  # yellow
    "#A142F4",  # purple
    "#FF6D00",  # orange
    "#24C1E0",  # cyan
    "#F439A0",  # pink
    "#0F9D58",  # dark green
    "#795548",  # brown
    "#7CB342",  # light green
    "#00838F",  # teal
    "#C5221F",  # dark red
    "#9E9E9E",  # grey
]

TRANSPORT_LABELS = {
    "owntruck": "Own Truck",
    "transporter": "Transporter",
    "owntruck/transporter": "Own Truck + Transporter",
}
TRANSPORT_COLORS = {
    "Own Truck": "#34A853",
    "Transporter": "#4285F4",
    "Own Truck + Transporter": "#FBBC04",
    "(not specified)": "#9AA0A6",
}

# Turnover tiers per drop point (THB)
TIERS = [
    ("A", "A · 100K+", 100_000, None),
    ("B", "B · 30K–100K", 30_000, 100_000),
    ("C", "C · 10K–30K", 10_000, 30_000),
    ("D", "D · under 10K", None, 10_000),
]

REGIONS = {
    "BKK": "Bangkok & Metro",
    "C": "Central",
    "E": "Eastern",
    "W": "Western",
    "N": "Northern",
    "NE": "Northeastern",
    "S": "Southern",
}


def region_of(lat, lng):
    """Coarse region split from the coordinate — a grouping aid, not a real province boundary."""
    if 13.40 <= lat <= 14.30 and 100.20 <= lng <= 100.95:
        return REGIONS["BKK"]
    if lat < 11.0:
        return REGIONS["S"]
    if lat > 16.5 and lng < 101.5:
        return REGIONS["N"]
    if lng > 101.5 and lat > 14.2:
        return REGIONS["NE"]
    if lng > 100.9 and lat <= 14.2:
        return REGIONS["E"]
    if lng < 100.2 and lat <= 16.5:
        return REGIONS["W"]
    return REGIONS["C"]


def tier_of(value):
    for _key, label, lo, hi in TIERS:
        if (lo is None or value >= lo) and (hi is None or value < hi):
            return label
    return TIERS[-1][1]


def clean(text):
    """Collapse whitespace and normalise text coming out of Excel."""
    if text is None:
        return ""
    return re.sub(r"\s+", " ", unicodedata.normalize("NFC", str(text))).strip()


def transport_of(raw):
    return TRANSPORT_LABELS.get(clean(raw).lower(), clean(raw) or "(not specified)")


def read_month(path):
    """Return (rows, skipped, notes) — skipped = rows without a usable coordinate."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Customers with Coordinates"]

    # The "No Match" sheet is left over from an earlier matching pass: every row in it
    # now carries a coordinate on the main sheet. Counting it again would double-count
    # the turnover, so instead flag those rows for a visual check.
    flagged = set()
    if "No Match" in wb.sheetnames:
        for raw in wb["No Match"].iter_rows(min_row=2, values_only=True):
            chain, code, name, _value = (list(raw) + [None] * 4)[:4]
            if clean(chain):
                flagged.add((clean(code).upper(), clean(name).upper()))

    rows, skipped = [], []
    for raw in ws.iter_rows(min_row=2, values_only=True):
        chain, code, name, value, orders, mode, lat, lng = (list(raw) + [None] * 8)[:8]
        chain = clean(chain)
        # trailing summary line, not a customer
        if not chain or chain.upper() == "TOTAL":
            continue
        value = float(value or 0)
        entry = {"chain": chain, "code": clean(code), "ship": clean(name), "value": round(value, 2)}
        try:
            lat, lng = float(lat), float(lng)
        except (TypeError, ValueError):
            skipped.append(entry)
            continue
        # a coordinate outside Thailand is bad data — keep it off the map
        if not (5.0 <= lat <= 21.0 and 96.0 <= lng <= 106.5):
            skipped.append(entry)
            continue
        entry.update({
            "orders": int(orders or 0),
            "transport": transport_of(mode),
            "region": region_of(lat, lng),
            "tier": tier_of(value),
            "lat": round(lat, 6),
            "lon": round(lng, 6),
            "check": 1 if (clean(code).upper(), clean(name).upper()) in flagged else 0,
        })
        rows.append(entry)

    skipped.sort(key=lambda r: -r["value"])
    notes = []
    if "Notes" in wb.sheetnames:
        notes = [clean(r[0]) for r in wb["Notes"].iter_rows(values_only=True) if clean(r[0])]
    return rows, skipped, notes


def build(src):
    rows, skipped, notes = read_month(src["xlsx"])

    totals = {}
    for r in rows:
        totals[r["chain"]] = totals.get(r["chain"], 0) + r["value"]
    order = sorted(totals, key=lambda c: (-totals[c], c))
    colors = {c: GOOGLE_PALETTE[i % len(GOOGLE_PALETTE)] for i, c in enumerate(order)}

    flagged = [r for r in rows if r["check"]]
    month_value = sum(r["value"] for r in rows) + sum(r["value"] for r in skipped)
    data = {
        "slug": src["slug"],
        "month": src["month"],
        "source": src["xlsx"].name,
        "rows": rows,
        "colors": colors,
        "transportColors": TRANSPORT_COLORS,
        "skipped": skipped,
        "notes": notes,
        "monthValue": round(month_value, 2),
    }

    vendor_css = (VENDOR / "leaflet.css").read_text(encoding="utf-8")
    vendor_js = "\n;\n".join(
        (VENDOR / f).read_text(encoding="utf-8") for f in ("leaflet.js", "leaflet-heat.js")
    )

    html = TEMPLATE.read_text(encoding="utf-8")
    for token, value in {
        "__TITLE__": src["title"],
        "__MONTH__": src["month"],
        "__VENDOR_CSS__": vendor_css,
        "__VENDOR_JS__": vendor_js,
        "__DATA_JSON__": json.dumps(data, ensure_ascii=False, separators=(",", ":")),
    }.items():
        html = html.replace(token, value)

    OUT_DIR.mkdir(exist_ok=True)
    out = OUT_DIR / f"{src['slug']}_Customer_Map.html"
    out.write_text(html, encoding="utf-8")

    print(f"{out.relative_to(ROOT)}")
    print(f"   {len(rows):,} drop points · {len(flagged)} to verify · {len(skipped)} without coordinates · "
          f"{len(colors)} retail groups · value {month_value:,.2f} · {out.stat().st_size/1024:.0f} KB")
    return out


if __name__ == "__main__":
    for s in SOURCES:
        build(s)
